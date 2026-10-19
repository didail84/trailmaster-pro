"""
Script d'import ponctuel du programme Excel (musculation, mobilité,
proprioception, calendrier 6 mois) vers la base TrailMaster Pro.

Usage : python3 importer_programme_excel.py chemin_vers_fichier.xlsx athlete_id
"""

import sys
import json
import re
from datetime import datetime

import openpyxl

from data.database import connexion, creer_base_de_donnees


MAPPING_TYPE_SEANCE = {
    "Musculation A – bas du corps": "Musculation",
    "Musculation B – chaîne postérieure": "Musculation",
    "Musculation C – stabilité": "Musculation",
    "Vélo": "Vélo",
    "Tapis incliné": "Tapis incliné",
    "Running goudron": "Sortie longue",
    "Rando familiale": "Randonnée",
}


def _parser_duree_ou_distance(texte):
    """Retourne (mode, valeur) — mode 'duree' en minutes ou 'distance'
    en km, selon ce que le texte source exprime réellement (ex.
    '3–4 km' pour une rando n'est pas une durée)."""

    if not texte:
        return "duree", None

    texte = str(texte)
    nombres = [float(n.replace(",", ".")) for n in re.findall(r"\d+[.,]?\d*", texte)]
    if not nombres:
        return "duree", None

    valeur_moyenne = sum(nombres) / len(nombres)

    if "km" in texte.lower():
        return "distance", round(valeur_moyenne, 1)
    if re.search(r"\bh\b", texte.lower()) or "heure" in texte.lower():
        return "duree", round(valeur_moyenne * 60)
    return "duree", round(valeur_moyenne)


def _parser_date_calendrier(texte_date):
    """Les dates du calendrier sont au format 'Monday 07/09/2026'."""
    if not texte_date:
        return None
    m = re.search(r"(\d{2})/(\d{2})/(\d{4})", str(texte_date))
    if not m:
        return None
    jour, mois, annee = m.groups()
    return f"{annee}-{mois}-{jour}"


def importer(chemin_fichier, athlete_id):
    creer_base_de_donnees()
    wb = openpyxl.load_workbook(chemin_fichier, data_only=True)
    conn = connexion()
    cur = conn.cursor()

    # -----------------------------------------------------
    # 1) MUSCULATION : exercices + charges de départ
    # -----------------------------------------------------
    charges_depart = {}
    ws = wb["Charges musculation"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            charges_depart[row[0]] = row[1]

    cur.execute("DELETE FROM musculation_exercices WHERE athlete_id=?", (athlete_id,))
    ws = wb["Musculation"]
    nb_exercices = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        groupe, exercice = row[0], row[1]
        blocs = list(row[2:8]) + [None] * (6 - len(row[2:8]))
        cur.execute(
            """INSERT INTO musculation_exercices
               (athlete_id, groupe_seance, exercice, charge_depart_kg,
                bloc1, bloc2, bloc3, bloc4, bloc5, bloc6)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (athlete_id, groupe, exercice, charges_depart.get(exercice),
             *[str(b) if b is not None else None for b in blocs[:6]]),
        )
        nb_exercices += 1

    # -----------------------------------------------------
    # 2) MOBILITÉ & PROPRIOCEPTION
    # -----------------------------------------------------
    cur.execute("DELETE FROM mobilite_exercices WHERE athlete_id=?", (athlete_id,))
    nb_mobilite = 0

    ws = wb["Mobilité & étirements"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        exercice, moment, zone = row[0], row[1], row[2]
        blocs = list(row[3:9]) + [None] * (6 - len(row[3:9]))
        cur.execute(
            """INSERT INTO mobilite_exercices
               (athlete_id, categorie, exercice, moment_ou_objectif, zone,
                bloc1, bloc2, bloc3, bloc4, bloc5, bloc6, progression)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (athlete_id, "Mobilité", exercice, moment, zone,
             *[str(b) if b is not None else None for b in blocs[:6]], None),
        )
        nb_mobilite += 1

    ws = wb["Proprioception"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        exercice, objectif = row[0], row[1]
        blocs = list(row[2:8]) + [None] * (6 - len(row[2:8]))
        progression = row[8] if len(row) > 8 else None
        cur.execute(
            """INSERT INTO mobilite_exercices
               (athlete_id, categorie, exercice, moment_ou_objectif, zone,
                bloc1, bloc2, bloc3, bloc4, bloc5, bloc6, progression)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (athlete_id, "Proprioception", exercice, objectif, None,
             *[str(b) if b is not None else None for b in blocs[:6]], progression),
        )
        nb_mobilite += 1

    # -----------------------------------------------------
    # 3) PLANNING MOBILITÉ-PROPRIO (hebdomadaire)
    # -----------------------------------------------------
    cur.execute("DELETE FROM planning_mobilite WHERE athlete_id=?", (athlete_id,))
    ws = wb["Planning mobilité-proprio"]
    nb_planning = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        cur.execute(
            """INSERT INTO planning_mobilite
               (athlete_id, jour, duree, bloc, contenu, intensite, moment)
               VALUES (?,?,?,?,?,?,?)""",
            (athlete_id, row[0], row[1], row[2], row[3], row[4], row[5]),
        )
        nb_planning += 1

    # -----------------------------------------------------
    # 4) CALENDRIER 6 MOIS -> Plan d'entraînement + Séances structurées
    # -----------------------------------------------------
    ws = wb["Calendrier 6 mois"]
    lignes_calendrier = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]

    dates_valides = [_parser_date_calendrier(r[1]) for r in lignes_calendrier]
    dates_valides = [d for d in dates_valides if d]
    date_debut = min(dates_valides) if dates_valides else None
    date_fin = max(dates_valides) if dates_valides else None

    cur.execute(
        """INSERT INTO plans_entrainement
           (athlete_id, nom, distance_cible, niveau, nb_semaines, jours_par_semaine,
            allure_objectif, date_debut, date_course)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (athlete_id, "Programme Trail 6 mois (importé Excel)", "Trail", "Intermédiaire",
         24, 7, None, date_debut, date_fin),
    )
    plan_id = cur.lastrowid

    nb_seances = 0
    for row in lignes_calendrier:
        semaine, jour_texte, type_seance_brut, duree, contenu, objectif, intensite = row[:7]
        date_iso = _parser_date_calendrier(jour_texte)
        if not date_iso:
            continue

        type_seance = MAPPING_TYPE_SEANCE.get(type_seance_brut, "Autre")
        mode, valeur = _parser_duree_ou_distance(duree)

        etapes = [{
            "type": "Effort",
            "mode": mode,
            "valeur": valeur,
            "allure_cible": None,
            "fc_cible": None,
        }]

        notes = f"{contenu or ''}\nDurée prévue : {duree or '—'} | Objectif : {objectif or ''} | Intensité : {intensite or ''}".strip()

        cur.execute(
            """INSERT INTO seances_structurees
               (athlete_id, nom, type_seance, date_prevue, statut, etapes_json, notes)
               VALUES (?,?,?,?,?,?,?)""",
            (athlete_id, type_seance_brut, type_seance, date_iso, "Prévue",
             json.dumps(etapes), notes),
        )
        nb_seances += 1

    conn.commit()
    conn.close()

    return {
        "exercices_musculation": nb_exercices,
        "exercices_mobilite": nb_mobilite,
        "planning_mobilite": nb_planning,
        "seances_importees": nb_seances,
        "plan_id": plan_id,
        "date_debut": date_debut,
        "date_fin": date_fin,
    }


if __name__ == "__main__":
    chemin = sys.argv[1]
    athlete_id = int(sys.argv[2])
    resultat = importer(chemin, athlete_id)
    print(resultat)
